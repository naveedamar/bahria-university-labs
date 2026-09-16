import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def build_fds_docx():
    doc = docx.Document()

    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    PRIMARY = RGBColor(27, 54, 93)
    SECONDARY = RGBColor(90, 105, 120)
    DARK = RGBColor(33, 37, 41)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PRIMARY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Arial'
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = DARK
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        rb = p.add_run(bold_prefix)
        rb.font.name = 'Arial'
        rb.font.size = Pt(10)
        rb.font.bold = True
        rb.font.color.rgb = DARK
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.color.rgb = DARK

    def format_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_row = table.rows[0]
        for idx, h_text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.text = h_text
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, top=120, bottom=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row_data in enumerate(data):
            row = table.add_row()
            bg_color = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(val)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(9)
                    r.font.color.rgb = DARK

        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    rt = p_title.add_run("Functional Design Specification (FDS)\n")
    rt.font.name = 'Arial'
    rt.font.size = Pt(22)
    rt.font.bold = True
    rt.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("University Cafeteria Pre-Order System\nSoftware Engineering Course Project | 2026")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = SECONDARY

    # Section 1
    add_h1("1. Executive Summary & Problem Statement")
    add_h2("1.1 Project Background")
    add_p(
        "Campus dining facilities encounter severe operational bottlenecks during peak hours (11:30 AM to 2:00 PM). "
        "Traditional in-person ordering forces students to wait in two sequential lines: first to place and pay for orders, "
        "and second to wait for preparation. This leads to overcrowding, inaccurate fulfillment, missed academic sessions, "
        "and kitchen inefficiency."
    )
    add_h2("1.2 Proposed Solution")
    add_p(
        "The University Cafeteria Pre-Order System is a responsive web application designed to decouple menu browsing "
        "and order staging from the physical counter. Students browse dynamic catalogs, configure carts, and reserve designated "
        "collection windows from their devices. Cafeteria staff receive tickets through an active digital queue, prepare items in "
        "advance, and transition order statuses in real time."
    )
    add_h2("1.3 System Scope")
    add_bullet("In-Scope: ", "Authentication, role-based views (Student, Staff, Manager), live menu presentation, cart adjustments, scheduled pickup windows, unique order number generation, lifecycle tracking, kitchen stock toggling, and menu CRUD management.")
    add_bullet("Out-of-Scope (Phase 1): ", "Online payment gateway integration (counter cash settlement assumed), external delivery dispatch, GPS live tracking, and multi-campus branch synchronization.")

    # Section 2
    add_h1("2. Stakeholder Matrix")
    stake_headers = ["Stakeholder", "Category", "Operational Interest & Responsibilities"]
    stake_data = [
        ["Students", "Primary End-Users", "Authenticate, browse items, verify price/availability, schedule orders, and track preparation statuses."],
        ["Cafeteria Staff", "Operational Users", "Inspect incoming order tickets, update ticket statuses, and quickly mark sold-out items as unavailable."],
        ["Cafeteria Manager", "Administrative Users", "Update prices, descriptions, and catalog items; manage availability; and inspect sales/order logs."],
        ["University Admin", "Institutional Oversight", "Evaluate campus service efficiency, queue reduction, hygiene standards, and operational auditability."],
        ["System Admin", "Technical Support", "Maintain application server availability, secure password records, manage role allocations, and prevent downtime."],
        ["Food Suppliers", "External Beneficiaries", "Benefit indirectly from consolidated advance order volumes to streamline raw ingredient deliveries."]
    ]
    t_stake = doc.add_table(rows=1, cols=3)
    format_table(t_stake, [1.5, 1.6, 3.4], stake_headers, stake_data)

    # Section 3
    add_h1("3. Functional Requirements (FR)")
    add_h2("3.1 Authentication & Authorization")
    add_bullet("FR-01 (User Registration): ", "The system shall permit students to register with a username, valid university email, unique student ID, and secure password.")
    add_bullet("FR-02 (User Authentication): ", "The system shall authenticate registered users using hashed password verification and maintain session state across HTTP requests.")
    add_bullet("FR-03 (Role-Based Access Control): ", "The system shall enforce role segregation (Student, Staff, Manager), redirecting or aborting requests (HTTP 403) whenever an unauthorized role attempts access to a protected route.")
    add_bullet("FR-04 (Session Invalidation): ", "The system shall provide an explicit logout function that purges session variables and client-side authentication cookies.")

    add_h2("3.2 Student Menu & Ordering")
    add_bullet("FR-05 (Menu Presentation): ", "The system shall present an active list of food items, displaying item name, description, price (PKR), category, and availability status.")
    add_bullet("FR-06 (Inventory Shielding): ", "The system shall visually flag unavailable items and prohibit adding sold-out items to the shopping cart.")
    add_bullet("FR-07 (Cart Management): ", "The system shall allow students to add items, increment quantities, decrement quantities, and remove line items dynamically.")
    add_bullet("FR-08 (Financial Calculation): ", "The system shall calculate line-item subtotals and an accurate order total in real time without rounding discrepancies.")
    add_bullet("FR-09 (Pickup Slot Scheduling): ", "The system shall mandate selecting a valid pickup time window prior to order finalization.")
    add_bullet("FR-10 (Order Number Generation): ", "The system shall assign an uppercase alphanumeric identifier (e.g., ORD-XXXXXX) to every confirmed transaction.")
    add_bullet("FR-11 (Order Lifecycle Tracking): ", "The system shall display personal order records to the student with current status markers (Pending, Preparing, Ready, Completed).")

    add_h2("3.3 Kitchen & Staff Operations")
    add_bullet("FR-12 (Live Order Queue): ", "The system shall present kitchen staff with an active queue of pending and in-progress tickets sorted chronologically.")
    add_bullet("FR-13 (Order Lifecycle Advancement): ", "The system shall provide one-click status transitions across defined stages: Pending -> Preparing -> Ready -> Completed.")
    add_bullet("FR-14 (Real-Time Stock Toggle): ", "The system shall provide kitchen operators with a direct switch to toggle item availability (Available <-> Unavailable) without manager intervention.")

    add_h2("3.4 Manager Administration")
    add_bullet("FR-15 (Catalog Expansion): ", "The system shall provide a manager interface to create new food items with name, category, unit price, and descriptions.")
    add_bullet("FR-16 (Catalog Update): ", "The system shall permit editing existing food item details, including pricing corrections and description updates.")
    add_bullet("FR-17 (Safe Item Deletion): ", "The system shall prevent hard deletions of menu items tied to historical order records, falling back to an automatic availability flag deactivation.")
    add_bullet("FR-18 (System Audit Log): ", "The system shall grant managers view access to all system orders, listing user associations, pickup times, totals, and lifecycle outcomes.")

    # Section 4
    add_h1("4. Non-Functional Requirements (NFR)")
    add_bullet("NFR-01 (Performance & Latency): ", "Server response times for dynamic page renders must stay under 1.5 seconds under standard concurrent campus load.")
    add_bullet("NFR-02 (Availability): ", "The application core should remain available throughout standard cafeteria preparation and dining hours (08:00 to 18:00 PKT).")
    add_bullet("NFR-03 (Security & Cryptography): ", "Passwords must never be stored in plaintext; all credentials must be salted and hashed via PBKDF2/SHA-256 via werkzeug.security.")
    add_bullet("NFR-04 (Data Integrity & ACID Properties): ", "Multi-item order creations must be wrapped inside atomic database transactions. If an individual item validation fails, the entire transaction rollbacks.")
    add_bullet("NFR-05 (Usability & Responsiveness): ", "The interface must utilize responsive CSS grid/flexbox layouts (Bootstrap 5) to display consistently across desktop, tablet, and mobile screens.")
    add_bullet("NFR-06 (Maintainability): ", "The codebase must strictly follow the Model-View-Template (MVT/MVC) pattern, separating data persistence (models.py), presentation templates (/templates), and business routing logic (app.py).")

    # Section 5
    add_h1("5. Architectural Design & Database Schema")
    add_h2("5.1 System Architecture")
    add_p(
        "The application uses a three-tier client-server architectural model:\n"
        "1. Presentation Layer: Jinja2 templates styled with Bootstrap 5 and customized CSS.\n"
        "2. Application / Business Logic Layer: Python Flask application server managing session state, role validation decorators, and request dispatching.\n"
        "3. Data Persistence Layer: SQLite relational database managed via Flask-SQLAlchemy Object-Relational Mapper (ORM)."
    )

    add_h2("5.2 Relational Data Dictionary")
    add_p("Table: users", bold_prefix="")
    u_headers = ["Attribute", "Data Type", "Constraints", "Description"]
    u_data = [
        ["id", "INTEGER", "Primary Key, Auto", "Unique internal identifier for the user account."],
        ["username", "VARCHAR(80)", "NOT NULL, UNIQUE", "Distinct system login name."],
        ["email", "VARCHAR(120)", "NOT NULL, UNIQUE", "University email address for contact and audit."],
        ["student_id", "VARCHAR(50)", "NULLABLE", "Institutional registration roll number (students only)."],
        ["password_hash", "VARCHAR(255)", "NOT NULL", "Salted cryptographic hash of the user password."],
        ["role", "VARCHAR(20)", "NOT NULL, DEFAULT 'student'", "Security group: 'student', 'staff', or 'manager'."],
        ["created_at", "DATETIME", "NOT NULL, DEFAULT UTC", "Account creation timestamp."]
    ]
    t_u = doc.add_table(rows=1, cols=4)
    format_table(t_u, [1.2, 1.3, 1.8, 2.2], u_headers, u_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_p("Table: menu_items", bold_prefix="")
    m_headers = ["Attribute", "Data Type", "Constraints", "Description"]
    m_data = [
        ["id", "INTEGER", "Primary Key, Auto", "Unique internal identifier for the food item."],
        ["name", "VARCHAR(100)", "NOT NULL", "Commercial display title of the dish."],
        ["description", "TEXT", "NOT NULL", "Detailed ingredients, allergen notes, or preparation info."],
        ["category", "VARCHAR(50)", "NOT NULL, DEFAULT 'Meals'", "Menu division: 'Meals', 'Snacks', or 'Beverages'."],
        ["price", "FLOAT", "NOT NULL", "Unit price denominated in PKR."],
        ["is_available", "BOOLEAN", "NOT NULL, DEFAULT True", "Global availability switch for kitchen stock management."]
    ]
    t_m = doc.add_table(rows=1, cols=4)
    format_table(t_m, [1.2, 1.3, 1.8, 2.2], m_headers, m_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_p("Table: orders", bold_prefix="")
    o_headers = ["Attribute", "Data Type", "Constraints", "Description"]
    o_data = [
        ["id", "INTEGER", "Primary Key, Auto", "Unique internal tracking integer."],
        ["order_number", "VARCHAR(30)", "NOT NULL, UNIQUE", "Public alphanumeric identifier (e.g., ORD-A1B2C3)."],
        ["user_id", "INTEGER", "NOT NULL, FK(users.id)", "Foreign key linking the order to the placing student."],
        ["pickup_time", "VARCHAR(50)", "NOT NULL", "Scheduled collection window selected by the student."],
        ["status", "VARCHAR(20)", "NOT NULL, DEFAULT 'Pending'", "Current status ('Pending', 'Preparing', 'Ready', 'Completed')."],
        ["total", "FLOAT", "NOT NULL, DEFAULT 0.0", "Calculated aggregate financial total of all items."],
        ["created_at", "DATETIME", "NOT NULL, DEFAULT UTC", "Order submission timestamp."]
    ]
    t_o = doc.add_table(rows=1, cols=4)
    format_table(t_o, [1.2, 1.3, 1.8, 2.2], o_headers, o_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_p("Table: order_items", bold_prefix="")
    oi_headers = ["Attribute", "Data Type", "Constraints", "Description"]
    oi_data = [
        ["id", "INTEGER", "Primary Key, Auto", "Unique line-item association identifier."],
        ["order_id", "INTEGER", "NOT NULL, FK(orders.id)", "Cascade-linked reference to the parent order record."],
        ["menu_item_id", "INTEGER", "NOT NULL, FK(menu_items.id)", "Foreign key linking to the purchased menu item."],
        ["quantity", "INTEGER", "NOT NULL, DEFAULT 1", "Total quantity of this specific item requested."],
        ["price", "FLOAT", "NOT NULL", "Frozen historical unit price at the time of purchase."]
    ]
    t_oi = doc.add_table(rows=1, cols=4)
    format_table(t_oi, [1.2, 1.3, 1.8, 2.2], oi_headers, oi_data)

    # Section 6
    add_h1("6. End-to-End User Operational Flows")
    add_h2("6.1 Student Flow")
    p_sf = doc.add_paragraph()
    p_sf.paragraph_format.space_after = Pt(6)
    r_sf = p_sf.add_run(
        "[Landing / Login]\n"
        "       │\n"
        "       ▼\n"
        "[Browse Live Menu] ──── (Item Out of Stock?) ──► [Disabled Button / Notice]\n"
        "       │ (Available)\n"
        "       ▼\n"
        "[Add Item(s) to Cart]\n"
        "       │\n"
        "       ▼\n"
        "[View Cart & Adjust Quantities]\n"
        "       │\n"
        "       ▼\n"
        "[Select Pickup Time Slot]\n"
        "       │\n"
        "       ▼\n"
        "[Confirm & Submit Order]\n"
        "       │\n"
        "       ▼\n"
        "[Server Validates Stock & Flushes Cart]\n"
        "       │\n"
        "       ▼\n"
        "[Receive Order Ref (e.g., ORD-7F3A9C)]\n"
        "       │\n"
        "       ▼\n"
        "[Track Order on 'My Orders' Dashboard] ──► [Collect Food at Counter]"
    )
    r_sf.font.name = 'Courier New'
    r_sf.font.size = Pt(8.5)

    add_h2("6.2 Kitchen Staff Flow")
    p_kf = doc.add_paragraph()
    p_kf.paragraph_format.space_after = Pt(6)
    r_kf = p_kf.add_run(
        "[Staff Authentication]\n"
        "       │\n"
        "       ▼\n"
        "[Inspect Kitchen Live Queue] ──► [Pending Orders Grouped by Pickup Slot]\n"
        "       │\n"
        "       ├─► (Physical Stock Depleted?) ──► [Toggle Item to 'Unavailable']\n"
        "       │\n"
        "       ├─► [Click: Mark 'Preparing']\n"
        "       │\n"
        "       ├─► [Click: Mark 'Ready'] ────────► [Student Notified on Screen]\n"
        "       │\n"
        "       └─► [Click: Mark 'Completed'] ────► [Order Archived]"
    )
    r_kf.font.name = 'Courier New'
    r_kf.font.size = Pt(8.5)

    # Section 7
    add_h1("7. Risk Analysis & Engineering Mitigations")
    risk_headers = ["Identified Risk", "Severity", "Potential Impact", "Engineering Mitigation Strategy"]
    risk_data = [
        ["Peak-Hour Concurrency Surges", "High", "Database locking, duplicate orders, or dropped requests during rush hours.", "Maintain shopping cart state in lightweight client-side sessions; use atomic database writes only at final checkout confirmation."],
        ["Out-of-Stock Ordering Glitches", "Medium", "Student orders an item that ran out moments earlier in the physical kitchen.", "Backend availability check during checkout commit; rolls back order with an explicit flash alert if an item is sold out."],
        ["Double Form Submissions", "Medium", "Multiple identical orders submitted if a student repeatedly clicks submit on a slow connection.", "Generate a unique UUID order reference server-side, commit the record, and immediately clear session['cart']."],
        ["Privilege Escalation", "High", "Students manually accessing /staff/dashboard or /manager/dashboard via URL manipulation.", "Centralized @role_required decorator verifying session['user_role'] server-side on every restricted route."],
        ["Data Orphanage on Menu Deletions", "Medium", "Manager deletes a food item that is referenced in historical sales records.", "Enforce foreign key relational safety; block hard deletions if past order line items exist, prompting an availability toggle instead."]
    ]
    t_risk = doc.add_table(rows=1, cols=4)
    format_table(t_risk, [1.4, 0.8, 1.9, 2.4], risk_headers, risk_data)

    output_path = "University_Cafeteria_FDS_Documentation.docx"
    doc.save(output_path)
    print(f"Generated: {output_path}")

if __name__ == "__main__":
    build_fds_docx()